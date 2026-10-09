from pathlib import Path


class UnsupportedSourceFormatError(ValueError):
    """Raised when no extractor is currently available for a file format."""


def _clean_text(text):
    return (text or "").strip()


def _extract_text_file(file_path):
    path = Path(file_path)
    text = path.read_text(encoding="utf-8", errors="replace")
    return [text] if _clean_text(text) else []


def _extract_powerpoint(file_path):
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError(
            "PowerPoint extraction requires the python-pptx package."
        ) from exc

    presentation = Presentation(str(file_path))
    chunks = []

    for slide in presentation.slides:
        parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = _clean_text(shape.text)
                if text:
                    parts.append(text)

        slide_text = "\n".join(parts).strip()

        if slide_text:
            chunks.append(slide_text)

    return chunks


def _extract_pdf(file_path):
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError(
            "PDF extraction requires the pypdf package."
        ) from exc

    reader = PdfReader(str(file_path))
    chunks = []

    for page in reader.pages:
        text = _clean_text(page.extract_text())
        if text:
            chunks.append(text)

    return chunks


def _extract_docx(file_path):
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise RuntimeError(
            "Word extraction requires the python-docx package."
        ) from exc

    document = Document(str(file_path))
    chunks = []
    current_parts = []

    def flush_current():
        content = "\n".join(
            part for part in current_parts if _clean_text(part)
        ).strip()
        if content:
            chunks.append(content)
        current_parts.clear()

    # Walk Word body elements in their original order so paragraphs and tables
    # remain in the same sequence as the source document.
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            paragraph_text = _clean_text(paragraph.text)
            if not paragraph_text:
                continue

            style_name = ""
            if paragraph.style is not None:
                style_name = _clean_text(paragraph.style.name).casefold()

            # Word's own heading structure defines section boundaries.
            # No business/domain vocabulary is hard-coded.
            if style_name.startswith("heading"):
                flush_current()
                current_parts.append(paragraph_text)
            else:
                current_parts.append(paragraph_text)

        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            table_rows = []

            for row in table.rows:
                values = []
                for cell in row.cells:
                    value = _clean_text(cell.text)
                    if value:
                        values.append(value)

                if values:
                    table_rows.append(" | ".join(values))

            if table_rows:
                # Preserve the table as its own searchable content unit while
                # keeping the surrounding Word structure intact.
                flush_current()
                chunks.append("\n".join(table_rows))

    flush_current()

    return chunks


def _extract_xlsx(file_path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Excel extraction requires the openpyxl package."
        ) from exc

    workbook = load_workbook(
        filename=str(file_path),
        read_only=True,
        data_only=True,
    )

    chunks = []

    try:
        for worksheet in workbook.worksheets:
            rows = []

            for row in worksheet.iter_rows(values_only=True):
                values = [
                    str(value).strip()
                    for value in row
                    if value is not None and str(value).strip()
                ]

                if values:
                    rows.append(" | ".join(values))

            if rows:
                chunks.append(
                    f"Worksheet: {worksheet.title}\n" + "\n".join(rows)
                )
    finally:
        workbook.close()

    return chunks


EXTRACTORS = {
    ".txt": _extract_text_file,
    ".md": _extract_text_file,
    ".csv": _extract_text_file,
    ".html": _extract_text_file,
    ".htm": _extract_text_file,
    ".pptx": _extract_powerpoint,
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
}


def supported_extensions():
    """
    Return the file extensions currently supported by the extraction layer.
    """
    return sorted(EXTRACTORS.keys())


def extract_content(file_path):
    """
    Extract text from a supported file without writing to the database.

    This layer only extracts content. It does not:
    - create KnowledgeSource records
    - create KnowledgeChunk records
    - generate embeddings
    - call an analysis model

    Unsupported formats fail explicitly rather than pretending they were read.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"Source file does not exist: {path}")

    if not path.is_file():
        raise ValueError(f"Source path is not a file: {path}")

    extension = path.suffix.lower()
    extractor = EXTRACTORS.get(extension)

    if extractor is None:
        raise UnsupportedSourceFormatError(
            f"Unsupported source format: {extension or '[no extension]'}. "
            f"Currently supported: {', '.join(supported_extensions())}"
        )

    chunks = extractor(path)

    return {
        "file_name": path.name,
        "extension": extension,
        "chunks": chunks,
        "chunk_count": len(chunks),
    }
